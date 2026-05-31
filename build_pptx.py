#!/usr/bin/env python3
"""Generate PrivacyNudging_Presentation.pptx from the presentation notes.
Dark theme; concise slides; full timed speaking script lives in the speaker notes."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- palette ----------
BG      = RGBColor(0x0E, 0x11, 0x16)
PANEL   = RGBColor(0x16, 0x1B, 0x26)
PANEL2  = RGBColor(0x1E, 0x25, 0x33)
ACCENT  = RGBColor(0x7C, 0x6C, 0xFF)
ACCENT2 = RGBColor(0x22, 0xD3, 0xEE)
TEXT    = RGBColor(0xE7, 0xEC, 0xF3)
DIM     = RGBColor(0xB6, 0xBF, 0xCF)
MUTE    = RGBColor(0x77, 0x82, 0x95)
GOOD    = RGBColor(0x34, 0xD3, 0x99)
WARN    = RGBColor(0xFB, 0xBF, 0x24)

EW, EH = Inches(13.333), Inches(7.5)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = EW
prs.slide_height = EH
BLANK = prs.slide_layouts[6]

slide_no = [0]

def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()

def base_slide(footer=True, divider=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    # left accent bar
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.16), EH)
    _solid(bar, ACCENT)
    if footer and not divider:
        slide_no[0] += 1
        f = s.shapes.add_textbox(Inches(0.5), Inches(7.02), Inches(10), Inches(0.4))
        tf = f.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = "Almuhimedi et al. (CHI 2015) — Mobile App Privacy Nudging"
        r.font.size = Pt(9); r.font.color.rgb = MUTE; r.font.name = FONT
        n = s.shapes.add_textbox(Inches(12.4), Inches(7.02), Inches(0.7), Inches(0.4))
        np = n.text_frame.paragraphs[0]; np.alignment = PP_ALIGN.RIGHT
        nr = np.add_run(); nr.text = str(slide_no[0])
        nr.font.size = Pt(9); nr.font.color.rgb = MUTE; nr.font.name = FONT
    return s

def add_runs(p, text, size, color, bold=False):
    """Parse **bold** segments into runs; bold parts use accent color."""
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if seg == "":
            continue
        r = p.add_run(); r.text = seg
        r.font.size = Pt(size); r.font.name = FONT
        is_bold = (i % 2 == 1)
        r.font.bold = bold or is_bold
        r.font.color.rgb = ACCENT2 if is_bold else color

def title(s, text, sub=None, top=0.45):
    tb = s.shapes.add_textbox(Inches(0.55), Inches(top), Inches(12.3), Inches(1.1))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = TEXT; r.font.name = FONT
    # accent underline
    ul = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(top + 0.78), Inches(1.2), Pt(4))
    _solid(ul, ACCENT)
    if sub:
        sb = s.shapes.add_textbox(Inches(0.58), Inches(top + 0.86), Inches(12), Inches(0.5))
        sp = sb.text_frame.paragraphs[0]
        sr = sp.add_run(); sr.text = sub
        sr.font.size = Pt(14); sr.font.italic = True; sr.font.color.rgb = MUTE; sr.font.name = FONT

def bullets(s, items, top=1.7, left=0.7, width=12.0, height=5.0, size=18, gap=10):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for it in items:
        if isinstance(it, tuple):
            text, level = it
        else:
            text, level = it, 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap); p.space_before = Pt(0)
        marker = "▸  " if level == 0 else "–  "
        mr = p.add_run(); mr.text = marker
        mr.font.size = Pt(size); mr.font.name = FONT
        mr.font.color.rgb = ACCENT if level == 0 else MUTE
        p.level = level
        if level > 0:
            pPr = p._pPr if p._pPr is not None else p.get_or_add_pPr()
            pPr.set('marL', str(Emu(Inches(0.4)).emu if hasattr(Emu(Inches(0.4)), 'emu') else 365760))
        add_runs(p, text, size if level == 0 else size - 2, TEXT if level == 0 else DIM)
    return tb

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

def panel(s, left, top, width, height, color=PANEL):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    _solid(box, color)
    return box

# =====================================================================
# 1. TITLE
# =====================================================================
s = base_slide(footer=False)
# big tag
tag = s.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(11), Inches(0.5))
tp = tag.text_frame.paragraphs[0]
tr = tp.add_run(); tr.text = "MOBILE HCI · ORAL EXAM"
tr.font.size = Pt(15); tr.font.bold = True; tr.font.color.rgb = ACCENT2; tr.font.name = FONT
# title
tb = s.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(12), Inches(2.3))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Your Location Has Been Shared 5,398 Times!"
r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = TEXT; r.font.name = FONT
p2 = tf.add_paragraph()
r2 = p2.add_run(); r2.text = "A Field Study on Mobile App Privacy Nudging"
r2.font.size = Pt(24); r2.font.color.rgb = DIM; r2.font.name = FONT
# citation
cb = s.shapes.add_textbox(Inches(0.72), Inches(4.7), Inches(12), Inches(1))
cp = cb.text_frame.paragraphs[0]
cr = cp.add_run(); cr.text = "Almuhimedi, Schaub, Sadeh, Adjerid, Acquisti et al. — CHI 2015"
cr.font.size = Pt(18); cr.font.color.rgb = TEXT; cr.font.name = FONT
cp2 = cb.text_frame.add_paragraph()
cr2 = cp2.add_run(); cr2.text = "Lecture 8 — Online Privacy   ·   + my own idea: PrivacyMoment"
cr2.font.size = Pt(14); cr2.font.italic = True; cr2.font.color.rgb = MUTE; cr2.font.name = FONT
ul = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(4.55), Inches(2), Pt(4))
_solid(ul, ACCENT)
notes(s, "Open calmly. ~6 minutes total: ~3:45 on the paper, ~2:15 on my own idea. "
          "I walk the paper through nine steps like a story, then tell my own idea in the SAME nine steps so it mirrors the paper.")

# =====================================================================
# 2. ROADMAP
# =====================================================================
s = base_slide()
title(s, "How I'll walk through it", "Nine steps, like a story — then design ideas, then my own idea in the same nine steps")
bullets(s, [
    "**Goal** → **Worldview** → **Method** → **Lab/Field + Time**",
    "**Design details** → **Participants** → **Procedure**",
    "**Data** → **Analysis + Findings**",
    "Then: **design implications** + **contribution type**",
    "Then: **my idea (PrivacyMoment)** — told in the same 9 steps",
], top=1.9, size=22, gap=16)
notes(s, "This roadmap keeps the talk structured and easy to follow. The examiner can see I'm using the course's study-design vocabulary as a backbone.")

# =====================================================================
# DIVIDER: THE PAPER
# =====================================================================
def divider(text, sub, mins):
    s = base_slide(divider=True)
    panel(s, 0, 2.6, 13.333, 2.3, PANEL)
    tb = s.shapes.add_textbox(Inches(0.9), Inches(2.95), Inches(11.5), Inches(1.4))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = TEXT; r.font.name = FONT
    p2 = tf.add_paragraph()
    r2 = p2.add_run(); r2.text = sub + "    ·    " + mins
    r2.font.size = Pt(16); r2.font.color.rgb = ACCENT2; r2.font.name = FONT
    return s

s = divider("THE PAPER", "Almuhimedi et al.", "≈ 3:45")
notes(s, "Now the paper itself, through the nine steps.")

# =====================================================================
# STEP 1 — Goal
# =====================================================================
s = base_slide()
title(s, "1 · Goal & research questions", "What problem do they solve?")
bullets(s, [
    "Problem = **information asymmetry**: users don't know how often / which data apps take in the background",
    "Linked to the **privacy paradox**: what users *say* ≠ what their settings allow",
    "**RQ1** — does a detailed **permission manager** help people review & change permissions?",
    "**RQ2** — can **privacy nudges** make it work even better?",
    "A nudge is **soft-paternalistic** — it guides, never forces. “Keep sharing” always stays.",
], top=1.85, gap=12)
notes(s, "[0:00–0:35] This paper is by Almuhimedi and colleagues, from CHI 2015, from the Lecture 8 topic on online privacy. "
          "The problem is information asymmetry — one side knows much more than the other. Smartphone users do not know how often their apps take private data in the background, or which data they take. "
          "The authors link this to the privacy paradox — the gap between what users say they care about and what their settings actually allow. "
          "So their goal is two questions. RQ1: does a detailed permission manager really help people check and change app permissions? "
          "RQ2: can privacy nudges — small reminders that gently push you to act — make this work even better? "
          "A nudge is soft-paternalistic: it guides you, but never forces you. 'Keep sharing' is always still an option.")

# =====================================================================
# STEPS 2-4 — Worldview / Method / Setting+Time
# =====================================================================
s = base_slide()
title(s, "2–4 · Worldview, method, setting & time")
bullets(s, [
    "**Worldview:** mostly **positivism** (count reviews, %, a regression) + an **interpretivist** part (interviews for the *why*) → **mixed methods, leaning positivist**",
    "**Method:** a **within-subjects quasi-experiment** — everyone does every condition, but **no concurrent no-nudge control group**, so not a clean experiment",
    "**Setting:** a real-life **field study** on people's own Android phones → far toward **realism** → high **ecological validity**",
    "**Time:** **longitudinal** — **22 days** (far from a one-off)",
], top=1.85, gap=14)
notes(s, "[0:35–0:50] Their worldview is mostly positivism — there is one measurable truth you can count. They count reviews, percentages, and run a regression. But they add a real interpretativism part through interviews, to explain the why. So it is mixed methods, leaning positivist.\n\n"
          "[0:50–1:05] The method is a mixed-methods, within-subjects quasi-experiment — everyone does every condition, but it is not a fully clean experiment. There is no separate no-nudge control group running at the same time. So it is a quasi-experiment.\n\n"
          "[1:05–1:25] The setting is a real-life field study — the real world, not a lab. It ran on people's own Android phones, with their own apps, at real times of day. So it sits far toward the realism side of the control–realism scale, which the authors say raises ecological validity. For time, it is longitudinal: 22 days. So it is far from a one-off test.")

# =====================================================================
# STEP 5 — Design details (with phase visual)
# =====================================================================
s = base_slide()
title(s, "5 · Design — the three phases", "IV = phase (3 levels) · within-subjects")
# phase strip
phases = [("Phase 1\nBaseline", "7 days · silent logging", PANEL2),
          ("Phase 2\nAppOps manager", "7 days · limit each app+permission", PANEL2),
          ("Phase 3\nManager + 1 nudge/day", "8 days · personalised nudge", ACCENT)]
x = 0.7
for name, desc, col in phases:
    box = panel(s, x, 1.85, 3.9, 1.5, col)
    tf = box.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.15)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = name
    r.font.size = Pt(15); r.font.bold = True
    r.font.color.rgb = BG if col == ACCENT else TEXT; r.font.name = FONT
    p2 = tf.add_paragraph()
    r2 = p2.add_run(); r2.text = desc
    r2.font.size = Pt(11); r2.font.color.rgb = BG if col == ACCENT else DIM; r2.font.name = FONT
    x += 4.1
bullets(s, [
    "The **nudge** = personalised full-screen card from each user's own logs: “your location has been shared **5,398 times**”",
    "Three options = **enhanced active choice**: change settings · show me more · keep sharing",
    "**Randomised nudge-type order** to cut **order effects** — but **phase order is fixed → not full counterbalancing**",
    "**DVs:** *reviewing* (opening AppOps) and *adjusting* (limit / allow)",
], top=3.6, gap=11, size=17)
notes(s, "[1:25–1:55] Now the design. It is within-subjects — everyone goes through all three phases. The real independent variable is the phase, with three levels. "
          "Phase 1 was a seven-day silent baseline, before any treatment. Phase 2 was the permission manager only, for seven days: they took Android's hidden AppOps and showed it through their study launcher, so users could limit each app-and-permission pair. "
          "Phase 3 was the manager plus one nudge a day, for eight days. The nudge is the new part: a personalised, full-screen privacy nudge built from each user's own access logs. The big line is one number: 'your location has been shared 5,398 times.' "
          "It gives three options, using enhanced active choice — you must make a clear choice: change settings, show-me-more, or keep-sharing. "
          "They randomly mixed the order of the four nudge types to reduce order effects, but did NOT mix the order of the phases — so this is not full counterbalancing. "
          "The main dependent variables are: reviewing (opening AppOps) and adjusting (limiting or allowing).")

# =====================================================================
# STEPS 6-7 — Participants & Procedure
# =====================================================================
s = base_slide()
title(s, "6–7 · Participants & procedure")
bullets(s, [
    "**N = 23** · 65% female · median age **23** · recruited locally (Craigslist + university pool), in person",
    "Needed a narrow **Android 4.3–4.4.1** (for AppOps) and had to promise **not to update** — **3 removed** for upgrading",
    "**Procedure (in phase order):** consent + entry survey → 7 silent baseline days → 7 days AppOps available → 8 days with 1 nudge/day → exit survey + interviews",
    "System **sampled AppOps state every 5 minutes**",
], top=1.95, gap=14)
notes(s, "[1:55–2:10] There were 23 participants, 65% female, with a median age of 23. They were recruited locally, through Craigslist and a university pool, and had to come in person. They needed a narrow Android version, 4.3 to 4.4.1, for AppOps, and had to promise not to update. Three people were later removed for upgrading.\n\n"
          "[2:10–2:30] The procedure follows the three phases in order. People came in, gave consent, and did an entry survey. Then Phase 1: seven silent baseline days, just logging. Then Phase 2: seven days with AppOps available. Then Phase 3: eight days with one nudge a day, built from their own logs. At the end, an exit survey and interviews. The system sampled the AppOps state every five minutes.")

# =====================================================================
# STEPS 8-9 — Data, analysis & findings (merged, with quote)
# =====================================================================
s = base_slide()
title(s, "8–9 · Data, analysis & findings")
bullets(s, [
    "**Data spine = observed, behavioral** logs (what people *did*) + **self-reported, attitudinal** Likert comfort & **8 interviews** for the *why*",
    "Fixed rules: **behavioral ⇒ observed** · **attitudinal ⇒ self-reported**",
    "Nudge taps: **42% change · 25% keep · 33% ignored**",
    "**Analysis:** counts & % + one **random-effects regression**; qualitative = **thematic analysis**",
    "**Findings (1)** manager alone works — **22/23** checked, **65%** changed    **(2)** nudges add value — **95%** re-checked, **58%** limited *more*",
    "**(3) Access frequency** was the magic — regression linked frequency → choosing to change",
], top=1.5, gap=8, size=15)
# quote panel
q = panel(s, 0.7, 5.5, 11.9, 1.25, PANEL)
qb = q.text_frame; qb.word_wrap = True; qb.margin_left = Inches(0.25); qb.margin_top = Inches(0.1)
qp = qb.paragraphs[0]
qr = qp.add_run(); qr.text = "“4,182 times… it felt like I'm being followed by my own phone. It was scary.”"
qr.font.size = Pt(16); qr.font.italic = True; qr.font.color.rgb = TEXT; qr.font.name = FONT
qp2 = qb.add_paragraph()
qr2 = qp2.add_run(); qr2.text = "— P10  ·  phone-as-extension-of-self"
qr2.font.size = Pt(11); qr2.font.color.rgb = ACCENT2; qr2.font.name = FONT
notes(s, "[2:30–3:25] DATA: most of it is observed, behavioral system logs — what people actually did: five-minute logs of reviews and restrictions, and nudge taps (42% change, 25% keep, 33% ignored). "
          "There is also self-reported attitudinal data — a five-level Likert comfort scale and interviews with eight people. Behavioral data is always observed; attitudinal is always self-reported. So the spine is observed behaviour, with self-report to explain the why.\n\n"
          "ANALYSIS & FINDINGS: the analysis is mixed — counts and percentages plus one random-effects regression, and thematic analysis on the qualitative side. The findings come in three parts. "
          "One: the manager alone already works — 22 of 23 checked, 65% made changes. "
          "Two: nudges add real value on top — even after a week, 95% checked again and 58% limited access even more. "
          "Three: access frequency was the magic — all eight interviewees said the big number grabbed them. "
          "P10 said '4,182 times… it felt like I'm being followed by my own phone. It was scary.' That shows phone-as-extension-of-self, and the regression linked frequency to choosing to change settings.")

# =====================================================================
# Design implications + contribution
# =====================================================================
s = base_slide()
title(s, "Design implications + contribution type")
bullets(s, [
    "**Personalise** — stop nudging about apps the user already accepted",
    "Make it **“salient, sticky, but not annoying”**",
    "Main job = **feedback / visibility of system status** (Nielsen #1) — makes invisible background access visible",
    "Fixes a **discoverability** problem — **91%** had never used AppOps",
    "Contribution = mainly **empirical** (new findings) + a strong **artifact** (the built nudge)",
], top=1.95, gap=13)
notes(s, "[3:25–3:45] So the design ideas are: personalise — stop nudging about apps the user already accepted. Make it 'salient, sticky, but not annoying' — easy to notice, stays around, not irritating. "
          "The nudge's main job is feedback / visibility of system status — Nielsen #1, showing what the system is doing: it makes invisible background access visible. "
          "It also fixes a discoverability problem — 91% had never used AppOps. The contribution is mainly empirical — new findings from a study — with a strong artifact contribution in the nudge, a built thing.")

# =====================================================================
# DIVIDER: MY IDEA
# =====================================================================
s = divider("MY IDEA — PrivacyMoment", "Keep the nudge that works; fix WHEN it fires", "≈ 2:15")
notes(s, "Now my own idea, told in the same nine steps so it mirrors the paper.")

# =====================================================================
# The gap + goal
# =====================================================================
s = base_slide()
title(s, "The gap, and my goal", "The paper is blind to the context of use")
bullets(s, [
    "Nudge fired at a **random time (11am–8pm)** → **P4** always got it at work; **P19** got one **mid-run**, covering their running app",
    "Interviews: pressing **“keep sharing” was often just killing a badly-timed interruption** — not a real privacy choice",
    "**PrivacyMoment goal:** keep the frequency nudge that works, but fire it **only at a good moment** — so a dismissal means a *real* choice",
], top=1.95, gap=16, size=19)
notes(s, "[3:45–4:15] That leads to my own idea. It targets the paper's clearest weak point: it is blind to the context of use — the situation the user is in. "
          "The nudge fired at a random time between 11am and 8pm. So P4 always got it at work. And P19 got one mid-run, and it covered their running app. "
          "The key point: when people pressed 'keep sharing,' the interviews show it was often just getting rid of a badly-timed interruption, not a real privacy choice. "
          "My idea is PrivacyMoment. Its goal: keep the frequency nudge that works, but fire it only at a good moment, so a dismissal means a real choice — not a fix for bad timing.")

# =====================================================================
# PrivacyMoment design
# =====================================================================
s = base_slide()
title(s, "PrivacyMoment — the design")
bullets(s, [
    "Same **soft-paternalistic frequency nudge** — but respects **temporal & task context**",
    "**Senses the situation:** foreground app running? phone moving (accelerometer)? calendar busy?",
    "Fires **only when the nudge can BE the foreground task** → no interruption",
    "**Multimodal sensing** (accelerometer + app-state + calendar); handles **situational impairment** (running, driving)",
    "Adds **“remind me later”** (reschedules instead of a fake “keep sharing”), **personalisation**, and **feedforward** (shows what would break before you revoke)",
], top=1.9, gap=12, size=17)
notes(s, "[4:15–4:45] PrivacyMoment is the same soft-paternalistic frequency nudge. But it respects temporal and task context — the time, and what task the user is doing. "
          "It does not fire at a random clock time. It senses the situation: is a foreground app running? Is the phone moving, using the accelerometer? Is the calendar busy? "
          "It fires only when the nudge can be the foreground task, so it does not break into another task. "
          "It is multimodal on the sensing side: the accelerometer for motion, the app-foreground state, and the calendar. It handles situational impairment, like running or driving. "
          "I add a sticky 'remind me later,' so a bad moment reschedules the nudge instead of forcing a fake 'keep sharing.' I personalise: once you mark an app as 'I'm comfortable with this,' it drops out. "
          "And I add feedforward — each option shows what will happen: before you revoke, it shows what might break.")

# =====================================================================
# How I'd study it — table (9 steps)
# =====================================================================
s = base_slide()
title(s, "How I'd study it — same 9 steps")
rows = [
    ("Step", "My choice"),
    ("Goal", "Does context-aware timing CAUSE fewer bad-timing dismissals?"),
    ("Worldview / Method", "Positivist · field experiment with a clean control group"),
    ("Setting + Time", "Field (ecological validity) · longitudinal 3–4 wks (kills novelty effects)"),
    ("Design", "Between-subjects (you can't un-see your stats → avoids order/learning effects)"),
    ("IV / Hypotheses", "IV = timing: (A) random control vs (B) context-aware · H₁ lower dismissal (pre-registered, no HARKing)"),
    ("Participants", "Diverse sample for external validity · power calc ≈ N=157"),
    ("Data / Analysis", "Observed logs + Likert comfort · Mann-Whitney U + effect size r"),
]
nrows, ncols = len(rows), 2
tbl_shape = s.shapes.add_table(nrows, ncols, Inches(0.7), Inches(1.75), Inches(11.95), Inches(4.9))
table = tbl_shape.table
table.columns[0].width = Inches(3.0)
table.columns[1].width = Inches(8.95)
# strip default style banding by setting our own colors
for ri, (a, b) in enumerate(rows):
    for ci, val in enumerate((a, b)):
        cell = table.cell(ri, ci)
        cell.fill.solid()
        if ri == 0:
            cell.fill.fore_color.rgb = ACCENT
        else:
            cell.fill.fore_color.rgb = PANEL if ri % 2 else PANEL2
        cell.margin_left = Inches(0.12); cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = val
        r.font.name = FONT
        if ri == 0:
            r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = BG
        else:
            r.font.size = Pt(12.5)
            r.font.bold = (ci == 0)
            r.font.color.rgb = TEXT if ci == 0 else DIM
notes(s, "[4:45–5:55] To test it, I follow the same nine steps. Goal: does context-aware timing cause fewer bad-timing dismissals than random timing? "
          "Worldview: positivist — I want a measurable, generalisable comparison. Method: a field experiment with a clean control group. "
          "Setting and time: a field study on their own phones, because the whole point is real interruptions, so I want ecological validity; and longitudinal, about three to four weeks, to kill novelty effects. "
          "Design: between-subjects, not within-subjects — because once you have seen your privacy stats, you cannot un-see them, so within-subjects would carry order and learning effects. "
          "My IV is the timing strategy, with two levels: random-time (the control) versus context-aware PrivacyMoment. My directional hypothesis H1 is that context-aware timing gives a lower bad-timing dismissal rate; the null H0 is no difference, written before the data to avoid HARKing. "
          "My key DV is that dismissal rate; I also measure number of reviews and number of restrictive changes. "
          "Participants: a more mixed sample than N=23 students, for external validity, with a power calculation pointing to about N=157. "
          "Procedure: entry session and consent, a baseline week, then random assignment to group A or B, then the intervention weeks with one nudge a day on location, contacts, calendar, and call logs, then an exit survey and interview. "
          "Data: observed behavioral logs plus self-reported comfort, on three axes. Analysis: the non-parametric Mann-Whitney U test with an effect size. "
          "The trade-off is on purpose: I sit toward realism and accept weaker internal validity, because field confounds like the connectivity problem this paper hit do creep in. I do this to learn how the nudge truly works in real life. Thank you.")

# =====================================================================
# THANK YOU
# =====================================================================
s = base_slide(footer=False)
panel(s, 0, 2.7, 13.333, 2.1, PANEL)
tb = s.shapes.add_textbox(Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.6))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Thank you"
r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = TEXT; r.font.name = FONT
p2 = tf.add_paragraph()
r2 = p2.add_run(); r2.text = "Keep the nudge that works — fix when it fires.  Questions?"
r2.font.size = Pt(18); r2.font.color.rgb = ACCENT2; r2.font.name = FONT
notes(s, "Pause. Invite questions. Backup slides follow: classification cheat-sheet, weaknesses, contexts of use, privacy terminology, and Q&A one-liners.")

# =====================================================================
# DIVIDER: BACKUP
# =====================================================================
s = divider("BACKUP SLIDES", "For Q&A — not part of the timed talk", "appendix")
notes(s, "Use these only if asked.")

# =====================================================================
# Backup: classification cheat-sheet
# =====================================================================
s = base_slide()
title(s, "Paper at a glance — classification")
rows = [
    ("Dimension", "Classification"),
    ("Setting", "Field study, far toward realism (their own phones, in the wild)"),
    ("Design", "Within-subjects quasi-experiment · 3 phases · NO control group"),
    ("Duration", "Longitudinal, 22 days (7+7+8) → opens habituation/novelty question"),
    ("Order effects", "Nudge-TYPE order randomised; PHASE order fixed → not full counterbalancing"),
    ("Worldview", "Positivist-leaning mixed methods (counts, %, regression p<.05)"),
    ("Internal validity", "Medium — phases confounded with time; dose varied; 10-min attribution"),
    ("External validity", "Weak — N=23, 65% female, local, platform-locked (it's the sample, not the setting)"),
    ("Ecological validity", "Strong — main strength"),
    ("Contribution", "Mainly empirical + artifact (the nudge) — Wobbrock & Kientz"),
]
t = s.shapes.add_table(len(rows), 2, Inches(0.7), Inches(1.7), Inches(11.95), Inches(5.1)).table
t.columns[0].width = Inches(3.2); t.columns[1].width = Inches(8.75)
for ri, (a, b) in enumerate(rows):
    for ci, val in enumerate((a, b)):
        cell = t.cell(ri, ci); cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT if ri == 0 else (PANEL if ri % 2 else PANEL2)
        cell.margin_left = Inches(0.12); cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]; cell.text_frame.word_wrap = True
        r = p.add_run(); r.text = val; r.font.name = FONT
        if ri == 0:
            r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = BG
        else:
            r.font.size = Pt(11.5); r.font.bold = (ci == 0)
            r.font.color.rgb = TEXT if ci == 0 else DIM
notes(s, "Quick reference if asked to classify the study on any axis.")

# =====================================================================
# Backup: weaknesses
# =====================================================================
s = base_slide()
title(s, "What the paper is weak on", "Say these first if questioned")
bullets(s, [
    "**No no-nudge control group** → Phase-3 effect confounded with time (internal validity)",
    "**Carry-over:** authors admit Phase 3 behaviour depends on Phase 2",
    "**Small, lopsided, self-chosen, platform-locked sample** (N=23) → weak external validity",
    "**Cause guessed from a 10-minute window**, not logged directly",
    "**Habituation / novelty** not handled in 8 nudge days (clear day-10 drop)",
    "**Restriction ≠ real privacy** — 68% of nudge-triggered changes were off-type",
    "Mostly **descriptive stats**, one regression, **no effect sizes**, underpowered (need ≈N=157)",
], top=1.85, gap=9, size=16)
notes(s, "These are the strongest critiques, ordered by what I'd raise first. Each ties to a course term: internal/external/ecological validity, carry-over confound, construct validity, novelty/habituation, power and effect size.")

# =====================================================================
# Backup: contexts of use
# =====================================================================
s = base_slide()
title(s, "Contexts of use (5 ISO types)")
bullets(s, [
    "**Temporal** (HIGH) — random 11am–8pm timing caused P4/P19 dismissals; habituation over 22 days",
    "**Task** (HIGH) — privacy is a *secondary/background* task; the nudge interrupts the *primary* task (P19's run)",
    "**Physical** (MED/HIGH) — body & movement: P19 getting a nudge **mid-run**",
    "**Technical** (HIGH) — Android 4.3–4.4 only; connectivity loss changed the dose; AppOps showed all permissions at once → 68% off-type",
    "**Social** (MED/HIGH) — privacy is social; user-vs-provider asymmetry; phone as personal territory",
], top=1.9, gap=11, size=16)
notes(s, "The paper is strong on technical and social context, weak on temporal and task — which is exactly the gap PrivacyMoment fixes. Keep ISO 'physical context' (body/movement) separate from the screen-space mobile design challenge.")

# =====================================================================
# Backup: privacy terminology chain
# =====================================================================
s = base_slide()
title(s, "Privacy concepts — the chain")
bullets(s, [
    "**Privacy paradox** — say ≠ do",
    "**Learned helplessness / privacy fatigue** — repeated invasion + no recourse → give up (explains the day-10 drop)",
    "**Information asymmetry** — providers know far more; the frequency number narrows it",
    "**Altman's regulation** — privacy is a *process* of adjusting access (a door, not a wall)",
    "**Nudge → soft-paternalism → enhanced active choice** — push, but keep the choice; warn against **dark patterns**",
    "**Leakiness** (data leaks out) vs **creepiness** (the felt intrusion) — Shklovski",
], top=1.85, gap=10, size=16)
notes(s, "Chain to say: concern exists (paradox) → invasion with no recourse → learned helplessness → people stop responding → a well-timed nudge restores control and cuts information asymmetry, which is Altman's idea of regulating access — as long as it never becomes a dark pattern.")

# =====================================================================
# Backup: Q&A one-liners
# =====================================================================
s = base_slide()
title(s, "Q&A quick-fire — 7 themes")
bullets(s, [
    "**Mobile challenges:** handles variable contexts, interruptions, discoverability; ignores task context & situational impairment",
    "**Design principles:** lead with **reification** (abstract → tappable card); main heuristic = visibility of system status; + feedforward",
    "**Modalities:** visual full-screen + single tap; fails under situational impairment; PrivacyMoment adds haptic + context sensing",
    "**Interaction qualities:** improves **discoverability**; neglects **expressiveness** (narrow tap input)",
    "**Context of use:** good on technical/social, bad on temporal/task → the PrivacyMoment gap",
    "**Data & methods:** mixed; observed-behavioral spine + attitudinal Likert; non-parametric correct",
    "**My study design:** longitudinal between-subjects field experiment; Mann-Whitney U + effect size",
], top=1.8, gap=8, size=14.5)
notes(s, "One line per course theme so I can answer any of the seven exam questions quickly, always linking the paper to my own idea.")

prs.save("/Users/todua/Documents/Github/Mobile HCI/PrivacyNudging_Presentation.pptx")
print("Saved PrivacyNudging_Presentation.pptx with", len(prs.slides._sldIdLst), "slides")
