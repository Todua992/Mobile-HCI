#!/usr/bin/env python3
"""Generic dark-theme deck renderer. Reads a deck JSON and writes a .pptx.
Usage: python build_deck.py <deck.json> <out.pptx>
Deck JSON shape: { paper{title,subtitle,authors,venue,topic}, talk[slide], designImplications slide,
  myIdea[slide], studyTable{title,rows,notes}, cheatSheet[[dim,cls]], weaknesses slide,
  contexts slide, terminology slide, qa slide }   where slide = {label,title,bullets[],notes,quote,quoteAttrib}
"""
import sys, json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------- palette ----------
BG      = RGBColor(0x0E, 0x11, 0x16)
PANEL   = RGBColor(0x16, 0x1B, 0x26)
PANEL2  = RGBColor(0x1E, 0x25, 0x33)
ACCENT  = RGBColor(0x7C, 0x6C, 0xFF)
ACCENT2 = RGBColor(0x22, 0xD3, 0xEE)
TEXT    = RGBColor(0xE7, 0xEC, 0xF3)
DIM     = RGBColor(0xB6, 0xBF, 0xCF)
MUTE    = RGBColor(0x77, 0x82, 0x95)
FONT = "Calibri"
EW, EH = Inches(13.333), Inches(7.5)

deck_path, out_path = sys.argv[1], sys.argv[2]
deck = json.load(open(deck_path))
P = deck["paper"]
FOOTER = f"{P['authors'].split(',')[0].split(' and ')[0].strip()} — {P['venue']}"

prs = Presentation(); prs.slide_width = EW; prs.slide_height = EH
BLANK = prs.slide_layouts[6]
pageno = [0]

def _solid(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()

def base_slide(footer=True, divider=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.16), EH); _solid(bar, ACCENT)
    if footer and not divider:
        pageno[0] += 1
        f = s.shapes.add_textbox(Inches(0.5), Inches(7.02), Inches(10), Inches(0.4))
        r = f.text_frame.paragraphs[0].add_run(); r.text = FOOTER
        r.font.size = Pt(9); r.font.color.rgb = MUTE; r.font.name = FONT
        n = s.shapes.add_textbox(Inches(12.4), Inches(7.02), Inches(0.7), Inches(0.4))
        np = n.text_frame.paragraphs[0]; np.alignment = PP_ALIGN.RIGHT
        nr = np.add_run(); nr.text = str(pageno[0]); nr.font.size = Pt(9); nr.font.color.rgb = MUTE; nr.font.name = FONT
    return s

def add_runs(p, text, size, color, bold=False):
    for i, seg in enumerate(text.split("**")):
        if seg == "": continue
        r = p.add_run(); r.text = seg; r.font.size = Pt(size); r.font.name = FONT
        is_b = (i % 2 == 1); r.font.bold = bold or is_b
        r.font.color.rgb = ACCENT2 if is_b else color

def title(s, text, sub=None, top=0.45):
    tb = s.shapes.add_textbox(Inches(0.55), Inches(top), Inches(12.3), Inches(1.1))
    tb.text_frame.word_wrap = True
    r = tb.text_frame.paragraphs[0].add_run(); r.text = text
    r.font.size = Pt(29); r.font.bold = True; r.font.color.rgb = TEXT; r.font.name = FONT
    ul = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(top + 0.78), Inches(1.2), Pt(4)); _solid(ul, ACCENT)
    if sub:
        sb = s.shapes.add_textbox(Inches(0.58), Inches(top + 0.86), Inches(12), Inches(0.5))
        sr = sb.text_frame.paragraphs[0].add_run(); sr.text = sub
        sr.font.size = Pt(14); sr.font.italic = True; sr.font.color.rgb = MUTE; sr.font.name = FONT

def bullets(s, items, top=1.7, left=0.7, width=12.0, size=18, gap=10):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.0))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        mr = p.add_run(); mr.text = "▸  "; mr.font.size = Pt(size); mr.font.name = FONT; mr.font.color.rgb = ACCENT
        add_runs(p, it, size, TEXT)

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text or ""

def panel(s, l, t, w, h, color=PANEL):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)); _solid(box, color); return box

def quote_panel(s, text, attrib, top=5.5, h=1.25):
    q = panel(s, 0.7, top, 11.9, h, PANEL)
    qb = q.text_frame; qb.word_wrap = True; qb.margin_left = Inches(0.25); qb.margin_top = Inches(0.1)
    qr = qb.paragraphs[0].add_run(); qr.text = "“" + text + "”"
    qr.font.size = Pt(16); qr.font.italic = True; qr.font.color.rgb = TEXT; qr.font.name = FONT
    if attrib:
        qp2 = qb.add_paragraph(); qr2 = qp2.add_run(); qr2.text = attrib
        qr2.font.size = Pt(11); qr2.font.color.rgb = ACCENT2; qr2.font.name = FONT

def fit_size(items):
    n = len(items); longest = max((len(x) for x in items), default=0)
    if n <= 4 and longest < 110: return 18, 12
    if n <= 5: return 16, 10
    return 15, 8

def content_slide(slide):
    s = base_slide()
    lab = (slide.get("label") or "").strip()
    ttl = (lab + " · " if lab else "") + slide["title"]
    title(s, ttl)
    has_quote = bool(slide.get("quote"))
    size, gap = fit_size(slide["bullets"])
    bullets(s, slide["bullets"], top=1.55, size=size, gap=gap)
    if has_quote:
        quote_panel(s, slide["quote"], slide.get("quoteAttrib", ""))
    notes(s, slide.get("notes", ""))

def table_slide(ttl, rows, header, note, col0=3.0):
    s = base_slide(); title(s, ttl)
    data = [header] + rows
    t = s.shapes.add_table(len(data), 2, Inches(0.7), Inches(1.7), Inches(11.95), Inches(5.0)).table
    t.columns[0].width = Inches(col0); t.columns[1].width = Inches(11.95 - col0)
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci); cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if ri == 0 else (PANEL if ri % 2 else PANEL2)
            cell.margin_left = Inches(0.12); cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            r = cell.text_frame.paragraphs[0].add_run(); r.text = val; r.font.name = FONT
            if ri == 0:
                r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = BG
            else:
                fs = 12.5 if len(data) <= 9 else 11.5
                r.font.size = Pt(fs); r.font.bold = (ci == 0); r.font.color.rgb = TEXT if ci == 0 else DIM
    notes(s, note)

def divider(text, sub):
    s = base_slide(divider=True)
    panel(s, 0, 2.6, 13.333, 2.3, PANEL)
    tb = s.shapes.add_textbox(Inches(0.9), Inches(2.95), Inches(11.5), Inches(1.4)); tb.text_frame.word_wrap = True
    r = tb.text_frame.paragraphs[0].add_run(); r.text = text
    r.font.size = Pt(38); r.font.bold = True; r.font.color.rgb = TEXT; r.font.name = FONT
    p2 = tb.text_frame.add_paragraph(); r2 = p2.add_run(); r2.text = sub
    r2.font.size = Pt(16); r2.font.color.rgb = ACCENT2; r2.font.name = FONT
    return s

# ============ TITLE ============
s = base_slide(footer=False)
r = s.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(11), Inches(0.5)).text_frame.paragraphs[0].add_run()
r.text = "MOBILE HCI · ORAL EXAM"; r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = ACCENT2; r.font.name = FONT
tb = s.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(12), Inches(2.6)); tb.text_frame.word_wrap = True
r1 = tb.text_frame.paragraphs[0].add_run(); r1.text = P["title"]
r1.font.size = Pt(38); r1.font.bold = True; r1.font.color.rgb = TEXT; r1.font.name = FONT
p2 = tb.text_frame.add_paragraph(); r2 = p2.add_run(); r2.text = P["subtitle"]
r2.font.size = Pt(20); r2.font.color.rgb = DIM; r2.font.name = FONT
ul = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(4.95), Inches(2), Pt(4)); _solid(ul, ACCENT)
cb = s.shapes.add_textbox(Inches(0.72), Inches(5.1), Inches(12), Inches(1.2))
cr = cb.text_frame.paragraphs[0].add_run(); cr.text = f"{P['authors']} — {P['venue']}"
cr.font.size = Pt(16); cr.font.color.rgb = TEXT; cr.font.name = FONT
cp2 = cb.text_frame.add_paragraph(); cr2 = cp2.add_run(); cr2.text = P["topic"]
cr2.font.size = Pt(14); cr2.font.italic = True; cr2.font.color.rgb = MUTE; cr2.font.name = FONT
notes(s, "Open calmly. ~6 minutes total: ~3:45 on the paper, ~2:15 on my own idea. I walk the paper through nine steps like a story, then tell my own idea in the same nine steps.")

# ============ ROADMAP ============
s = base_slide(); title(s, "How I'll walk through it", "Nine steps, like a story — then design ideas, then my own idea in the same steps")
bullets(s, [
    "**Goal** → **Worldview** → **Method** → **Setting + Time**",
    "**Design** → **Participants** → **Procedure**",
    "**Data** → **Analysis + Findings**",
    "Then: **design implications** + **contribution type**",
    "Then: **my own idea** — told in the same 9 steps",
], top=1.9, size=22, gap=16)
notes(s, "This roadmap keeps the talk structured. It shows the examiner I am using the course's study-design vocabulary as a backbone.")

# ============ THE PAPER ============
divider("THE PAPER", f"{P['authors'].split(',')[0].strip()} et al.    ·    ≈ 3:45")
for sl in deck["talk"]:
    content_slide(sl)
content_slide(deck["designImplications"])

# ============ MY IDEA ============
divider("MY IDEA", "Extending the paper    ·    ≈ 2:15")
for sl in deck["myIdea"]:
    content_slide(sl)
st = deck["studyTable"]
table_slide(st["title"], [list(r) for r in st["rows"]], ["Step", "My choice"], st["notes"], col0=3.0)

# ============ THANK YOU ============
s = base_slide(footer=False)
panel(s, 0, 2.7, 13.333, 2.1, PANEL)
tb = s.shapes.add_textbox(Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.6)); tb.text_frame.word_wrap = True
r = tb.text_frame.paragraphs[0].add_run(); r.text = "Thank you"
r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = TEXT; r.font.name = FONT
p2 = tb.text_frame.add_paragraph(); r2 = p2.add_run(); r2.text = "Questions?"
r2.font.size = Pt(18); r2.font.color.rgb = ACCENT2; r2.font.name = FONT
notes(s, "Pause. Invite questions. Backup slides follow: classification cheat-sheet, weaknesses, contexts of use, terminology, and Q&A one-liners.")

# ============ BACKUP ============
divider("BACKUP SLIDES", "For Q&A — not part of the timed talk")
table_slide("Paper at a glance — classification", [list(r) for r in deck["cheatSheet"]],
            ["Dimension", "Classification"], "Quick reference if asked to classify the study on any axis.", col0=3.2)
content_slide({**deck["weaknesses"], "label": "", "title": deck["weaknesses"].get("title", "What the paper is weak on")})
content_slide({**deck["contexts"], "label": "", "title": deck["contexts"].get("title", "Contexts of use")})
content_slide({**deck["terminology"], "label": "", "title": deck["terminology"].get("title", "Key concepts I must know")})
content_slide({**deck["qa"], "label": "", "title": deck["qa"].get("title", "Q&A quick-fire")})

prs.save(out_path)
print("Saved", out_path, "with", len(prs.slides._sldIdLst), "slides")
